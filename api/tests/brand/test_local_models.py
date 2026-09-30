"""OxeePhone lot 3: Local Models configuration, model listing, embeddings and
local document processing."""

import json

import httpx
import pytest

from api.brand import documents, embeddings, local_models
from api.schemas.ai_model_configuration import OrganizationAIModelConfigurationV2
from api.services.configuration.registry import (
    REGISTRY,
    ServiceType,
    SpeachesEmbeddingsConfiguration,
)


def _schemas(service_type: ServiceType) -> dict[str, dict]:
    return {
        provider: cls.model_json_schema()
        for provider, cls in REGISTRY[service_type].items()
    }


# ------------------------------------------------------------ configuration


@pytest.mark.parametrize(
    "service_type",
    [ServiceType.LLM, ServiceType.TTS, ServiceType.STT, ServiceType.EMBEDDINGS],
)
def test_only_local_models_are_offered(service_type):
    restricted = local_models.restrict_provider_schemas(
        service_type, _schemas(service_type)
    )
    assert list(restricted) == ["speaches"]
    schema = restricted["speaches"]
    assert schema["title"] == "Local Models"
    assert "provider_docs_url" not in schema
    assert schema["properties"]["base_url"]["default"] == ""
    assert schema["properties"]["model"]["default"] == ""
    assert "examples" not in schema["properties"]["model"]


def test_realtime_is_not_offered_and_stt_defaults_to_french():
    assert (
        local_models.restrict_provider_schemas(
            ServiceType.REALTIME, _schemas(ServiceType.REALTIME)
        )
        == {}
    )
    stt = local_models.restrict_provider_schemas(
        ServiceType.STT, _schemas(ServiceType.STT)
    )
    assert stt["speaches"]["properties"]["language"]["default"] == "fr"


def _pipeline(**providers):
    services = {
        "llm": {"provider": "speaches", "model": "m", "base_url": "http://x/v1"},
        "tts": {"provider": "speaches", "model": "t", "base_url": "http://x/v1"},
        "stt": {"provider": "speaches", "model": "s", "base_url": "http://x/v1"},
    }
    for service, provider in providers.items():
        services[service] = {
            "provider": provider,
            "model": "gpt-4.1",
            "api_key": "sk-test",
        }
    return OrganizationAIModelConfigurationV2.model_validate(
        {"version": 2, "mode": "byok", "byok": {"mode": "pipeline", "pipeline": services}}
    )


def test_local_models_pipeline_is_accepted():
    local_models.enforce_local_models(_pipeline())


def test_third_party_provider_is_rejected():
    with pytest.raises(ValueError) as exc:
        local_models.enforce_local_models(_pipeline(llm="openai"))
    assert exc.value.args[0] == [
        {"model": "llm", "message": "Only Local Models can be used."}
    ]


def test_dograh_managed_mode_is_rejected():
    configuration = OrganizationAIModelConfigurationV2.model_validate(
        {"version": 2, "mode": "dograh", "dograh": {"api_key": "key"}}
    )
    with pytest.raises(ValueError):
        local_models.enforce_local_models(configuration)


def test_masked_key_is_resolved_from_stored_configuration():
    stored = OrganizationAIModelConfigurationV2.model_validate(
        {
            "version": 2,
            "mode": "byok",
            "byok": {
                "mode": "pipeline",
                "pipeline": {
                    "llm": {
                        "provider": "speaches",
                        "model": "m",
                        "base_url": "http://x/v1",
                        "api_key": "secret-abcd",
                    },
                    "tts": {"provider": "speaches", "model": "t", "base_url": "http://x"},
                    "stt": {"provider": "speaches", "model": "s", "base_url": "http://x"},
                },
            },
        }
    )
    assert local_models.resolve_api_key("*******abcd", stored, "llm") == "secret-abcd"
    assert local_models.resolve_api_key("*******zzzz", stored, "llm") is None
    assert local_models.resolve_api_key("plain", stored, "llm") == "plain"
    assert local_models.resolve_api_key("", stored, "llm") is None


# ------------------------------------------------------------ model listing


@pytest.fixture
def mock_http(monkeypatch):
    """Route every httpx.AsyncClient created by the brand module to a handler."""
    state = {"handler": None, "requests": []}
    real_client = httpx.AsyncClient

    def factory(*args, **kwargs):
        def handle(request):
            state["requests"].append(request)
            return state["handler"](request)

        kwargs["transport"] = httpx.MockTransport(handle)
        return real_client(*args, **kwargs)

    monkeypatch.setattr(local_models.httpx, "AsyncClient", factory)
    return state


async def test_lists_models_from_openai_compatible_endpoint(mock_http):
    mock_http["handler"] = lambda request: httpx.Response(
        200, json={"object": "list", "data": [{"id": "b-model"}, {"id": "a-model"}]}
    )
    models = await local_models.list_endpoint_models("http://vllm:8000/v1/", "key")
    assert models == ["a-model", "b-model"]
    request = mock_http["requests"][0]
    assert str(request.url) == "http://vllm:8000/v1/models"
    assert request.headers["Authorization"] == "Bearer key"


async def test_listing_without_key_sends_no_authorization(mock_http):
    mock_http["handler"] = lambda request: httpx.Response(200, json={"data": []})
    await local_models.list_endpoint_models("http://vllm/v1", None)
    assert "Authorization" not in mock_http["requests"][0].headers


@pytest.mark.parametrize(
    "response, message",
    [
        (httpx.Response(401), "rejected the API key"),
        (httpx.Response(404), "HTTP 404"),
        (httpx.Response(200, text="<html>"), "did not return JSON"),
    ],
)
async def test_listing_errors_are_user_facing(mock_http, response, message):
    mock_http["handler"] = lambda request: response
    with pytest.raises(local_models.LocalModelsError, match=message):
        await local_models.list_endpoint_models("http://vllm/v1", None)


async def test_listing_requires_a_base_url():
    with pytest.raises(local_models.LocalModelsError, match="base URL"):
        await local_models.list_endpoint_models("  ", None)


# ------------------------------------------------------------ transcription


async def test_recording_transcription_uses_organization_stt(mock_http, monkeypatch):
    from api.schemas.ai_model_configuration import EffectiveAIModelConfiguration
    from api.services.configuration import ai_model_configuration

    effective = EffectiveAIModelConfiguration.model_validate(
        {
            "stt": {
                "provider": "speaches",
                "model": "whisper-large-v3",
                "base_url": "http://stt/v1",
                "api_key": "k",
            }
        }
    )

    class Resolved:
        pass

    resolved = Resolved()
    resolved.effective = effective

    async def fake_resolved(**_kwargs):
        return resolved

    monkeypatch.setattr(
        ai_model_configuration, "get_resolved_ai_model_configuration", fake_resolved
    )
    mock_http["handler"] = lambda request: httpx.Response(200, json={"text": " Bonjour "})

    result = await local_models.transcribe_with_organization_stt(
        organization_id=1,
        audio_data=b"RIFF",
        filename="a.wav",
        content_type="audio/wav",
        language="fr",
    )
    assert result == {"transcript": "Bonjour"}
    request = mock_http["requests"][0]
    assert str(request.url) == "http://stt/v1/audio/transcriptions"
    body = request.content.decode(errors="ignore")
    assert "whisper-large-v3" in body and 'name="language"' in body


# ------------------------------------------------------------ embeddings


def test_short_vectors_are_zero_padded_and_keep_cosine():
    vector = [0.6, 0.8]
    padded = embeddings.pad_embedding(vector, "m")
    assert len(padded) == 1536
    assert padded[:2] == vector and not any(padded[2:])


def test_oversized_vectors_are_rejected():
    with pytest.raises(embeddings.EmbeddingDimensionTooLargeError, match="4096"):
        embeddings.pad_embedding([0.0] * 4096, "big-model")


def test_speaches_embeddings_config_needs_no_key():
    config = SpeachesEmbeddingsConfiguration(model="bge-m3", base_url="http://e/v1")
    assert config.api_key is None


async def test_factory_builds_local_embedding_service():
    from api.services.gen_ai import build_embedding_service

    service = await build_embedding_service(
        db_client=None,
        provider="speaches",
        api_key=None,
        model="bge-m3",
        base_url="http://e/v1",
    )
    assert isinstance(service, embeddings.LocalEmbeddingService)
    assert service.get_model_id() == "bge-m3"


# ------------------------------------------------------------ documents


def test_markdown_is_chunked_with_heading_context(tmp_path):
    path = tmp_path / "guide.md"
    path.write_text(
        "# Accueil\n\nBonjour et bienvenue.\n\n## Horaires\n\nOuvert du lundi au vendredi.\n"
    )
    result = documents.process_document_locally(
        file_path=str(path), filename="guide.md", retrieval_mode="chunked", max_tokens=128
    )
    chunks = result["chunks"]
    assert [c["chunk_text"] for c in chunks] == [
        "Bonjour et bienvenue.",
        "Ouvert du lundi au vendredi.",
    ]
    assert chunks[1]["contextualized_text"] == (
        "Accueil > Horaires\nOuvert du lundi au vendredi."
    )
    assert [c["chunk_index"] for c in chunks] == [0, 1]
    assert result["docling_metadata"]["processor"] == "oxeephone-local"


def test_long_text_is_split_under_the_token_budget(tmp_path):
    path = tmp_path / "long.txt"
    sentence = "Ceci est une phrase de test assez longue pour remplir un morceau. "
    path.write_text(sentence * 100)
    result = documents.process_document_locally(
        file_path=str(path), filename="long.txt", retrieval_mode="chunked", max_tokens=32
    )
    assert len(result["chunks"]) > 1
    assert all(
        len(c["chunk_text"]) <= 32 * documents.CHARS_PER_TOKEN for c in result["chunks"]
    )


def test_txt_keeps_hash_lines_as_text(tmp_path):
    path = tmp_path / "notes.txt"
    path.write_text("# pas un titre\nligne suivante\n")
    result = documents.process_document_locally(
        file_path=str(path), filename="notes.txt", retrieval_mode="chunked", max_tokens=128
    )
    assert result["chunks"][0]["chunk_text"] == "# pas un titre\nligne suivante"


def test_docx_headings_and_tables(tmp_path):
    import docx

    document = docx.Document()
    document.add_heading("Tarifs", level=1)
    document.add_paragraph("Consultation standard.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Acte", "Prix"
    table.cell(1, 0).text, table.cell(1, 1).text = "Visite", "30 €"
    path = tmp_path / "tarifs.docx"
    document.save(path)

    result = documents.process_document_locally(
        file_path=str(path), filename="tarifs.docx", retrieval_mode="chunked", max_tokens=128
    )
    text = result["chunks"][0]["contextualized_text"]
    assert text.startswith("Tarifs\n")
    assert "Consultation standard." in text and "Visite | 30 €" in text


def test_pdf_text_is_extracted(tmp_path):
    from pypdf import PdfWriter
    from pypdf.generic import (
        DecodedStreamObject,
        DictionaryObject,
        NameObject,
    )

    writer = PdfWriter()
    page = writer.add_blank_page(width=300, height=200)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    page[NameObject("/Resources")] = DictionaryObject(
        {NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})}
    )
    stream = DecodedStreamObject()
    stream.set_data(b"BT /F1 12 Tf 20 100 Td (Horaires du cabinet) Tj ET")
    page[NameObject("/Contents")] = writer._add_object(stream)
    path = tmp_path / "doc.pdf"
    with open(path, "wb") as f:
        writer.write(f)

    result = documents.process_document_locally(
        file_path=str(path), filename="doc.pdf", retrieval_mode="chunked", max_tokens=128
    )
    assert result["chunks"][0]["chunk_text"] == "Horaires du cabinet"
    assert result["chunks"][0]["chunk_metadata"]["pages"] == [1]


def test_scanned_pdf_gets_a_clear_error(tmp_path):
    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    path = tmp_path / "scan.pdf"
    with open(path, "wb") as f:
        writer.write(f)
    with pytest.raises(documents.DocumentProcessingError, match="scanned"):
        documents.process_document_locally(
            file_path=str(path), filename="scan.pdf", retrieval_mode="chunked", max_tokens=128
        )


@pytest.mark.parametrize("name, message", [("old.doc", ".docx"), ("x.xlsx", "Unsupported")])
def test_unsupported_formats_get_a_clear_error(tmp_path, name, message):
    path = tmp_path / name
    path.write_bytes(b"data")
    with pytest.raises(documents.DocumentProcessingError, match=message):
        documents.process_document_locally(
            file_path=str(path), filename=name, retrieval_mode="chunked", max_tokens=128
        )


def test_full_document_mode_returns_text_only(tmp_path):
    path = tmp_path / "faq.json"
    path.write_text(json.dumps({"question": "Horaires ?", "reponse": "9h-18h"}))
    result = documents.process_document_locally(
        file_path=str(path), filename="faq.json", retrieval_mode="full_document", max_tokens=128
    )
    assert result["chunks"] == []
    assert '"reponse": "9h-18h"' in result["full_text"]
